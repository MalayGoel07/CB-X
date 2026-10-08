import os, re, uuid
from io import BytesIO
from pathlib import Path
from typing import Optional
from xml.sax.saxutils import escape

MIME = {
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}

_VERB = r"\b(make|create|generate|write|export|give|build|produce|save|convert|prepare|draft|turn)\b"
_NOT_CODE = r"\b(code|script|function|python|javascript|program|api|library|class|regex)\b"
_FORMATS = {"pdf":  r"\bpdf\b","docx": r"\b(docx?|word (doc|document|file)|ms word|microsoft word|as word|in word)\b","pptx": r"\b(pptx?|powerpoint|slides?|slide deck|presentation|deck)\b","xlsx": r"\b(xlsx?|excel|spreadsheet|workbook)\b",}

def detect_format(message: str) -> Optional[str]:
    m = message.lower()
    if not re.search(_VERB, m) or re.search(_NOT_CODE, m):
        return None
    for fmt, pat in _FORMATS.items():
        if re.search(pat, m):
            return fmt
    return None

_BASE = ("Output ONLY Markdown. No preamble, no closing remarks, no code fences around the whole thing. "
         "Use # for the title, ## for sections, - for bullets, and | tables | when data is tabular.")
FORMAT_RULES = {
    "pdf":  f"{_BASE} Write a well-structured document with a title and clear sections.",
    "docx": f"{_BASE} Write a well-structured document with a title and clear sections.",
    "pptx": f"{_BASE} Use one # line for the deck title, then one ## heading per slide with 3-5 short bullets each. "
            "Keep bullets under 12 words.",
    "xlsx": f"{_BASE} Output one or more Markdown tables. Put a ## heading before each table "
            "(it becomes the sheet name). Use plain numbers without currency symbols or commas in numeric cells.",
}

def clean_markdown(text: str) -> str:
    text = re.sub(r"<think>.*?</think>", "", text, flags=re.S).strip()
    m = re.fullmatch(r"```(?:markdown|md)?\s*\n(.*?)\n```", text, flags=re.S)
    return (m.group(1) if m else text).strip()

# ---------- markdown -> blocks ----------
def parse_markdown(md: str) -> list[tuple]:
    blocks, para, lines, i = [], [], md.splitlines(), 0

    def flush():
        if para:
            blocks.append(("p", " ".join(para)))
            para.clear()

    while i < len(lines):
        s = lines[i].strip()
        if s.startswith("```"):
            flush(); code = []; i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code.append(lines[i]); i += 1
            blocks.append(("code", "\n".join(code)))
        elif not s:
            flush()
        elif m := re.match(r"(#{1,6})\s+(.*)", s):
            flush(); blocks.append(("h", len(m[1]), m[2].strip()))
        elif s.startswith("|"):
            flush(); rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-+:?", c) for c in cells):  # skip |---|---| row
                    rows.append(cells)
                i += 1
            i -= 1
            w = max((len(r) for r in rows), default=0)
            blocks.append(("table", [r + [""] * (w - len(r)) for r in rows]))
        elif m := re.match(r"[-*•]\s+(.*)", s):
            flush(); blocks.append(("ul", m[1]))
        elif m := re.match(r"\d+[.)]\s+(.*)", s):
            flush(); blocks.append(("ol", m[1]))
        else:
            para.append(s)
        i += 1
    flush()
    return blocks

def _plain(t: str) -> str:
    return re.sub(r"\*\*(.+?)\*\*|`(.+?)`", lambda m: m.group(1) or m.group(2), t)

# ---------- renderers (each returns bytes) ----------
def render_docx(blocks, title) -> bytes:
    from docx import Document
    from docx.shared import Pt
    d = Document()

    def runs(par, text):
        for k, part in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
            if part:
                par.add_run(part).bold = bool(k % 2)

    for b in blocks:
        if b[0] == "h":
            d.add_heading(_plain(b[2]), level=0 if b[1] == 1 else min(b[1] - 1, 4))
        elif b[0] == "p":
            runs(d.add_paragraph(), b[1])
        elif b[0] == "ul":
            runs(d.add_paragraph(style="List Bullet"), b[1])
        elif b[0] == "ol":
            runs(d.add_paragraph(style="List Number"), b[1])
        elif b[0] == "code":
            r = d.add_paragraph().add_run(b[1]); r.font.name = "Consolas"; r.font.size = Pt(9)
        elif b[0] == "table" and b[1]:
            t = d.add_table(rows=len(b[1]), cols=len(b[1][0])); t.style = "Table Grid"
            for ri, row in enumerate(b[1]):
                for ci, c in enumerate(row):
                    cell = t.cell(ri, ci); cell.text = _plain(c)
                    if ri == 0:
                        for r in cell.paragraphs[0].runs: r.bold = True
            d.add_paragraph()
    buf = BytesIO(); d.save(buf); return buf.getvalue()

def render_pdf(blocks, title) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import Paragraph, Preformatted, SimpleDocTemplate, Spacer, Table, TableStyle

    st = getSampleStyleSheet()
    rl = lambda t: re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", escape(t))
    story, n = [], 0
    for b in blocks:
        n = n + 1 if b[0] == "ol" else 0
        if b[0] == "h":
            story += [Paragraph(rl(b[2]), st[f"Heading{min(b[1], 3)}"]), Spacer(1, 4)]
        elif b[0] == "p":
            story += [Paragraph(rl(b[1]), st["BodyText"]), Spacer(1, 6)]
        elif b[0] == "ul":
            story.append(Paragraph(rl(b[1]), st["BodyText"], bulletText="•"))
        elif b[0] == "ol":
            story.append(Paragraph(rl(b[1]), st["BodyText"], bulletText=f"{n}."))
        elif b[0] == "code":
            story += [Preformatted(b[1], st["Code"]), Spacer(1, 6)]
        elif b[0] == "table" and b[1]:
            data = [[Paragraph(rl(c), st["BodyText"]) for c in row] for row in b[1]]
            t = Table(data, repeatRows=1)
            t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                                   ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
                                   ("VALIGN", (0, 0), (-1, -1), "TOP")]))
            story += [t, Spacer(1, 8)]
    buf = BytesIO()
    SimpleDocTemplate(buf, pagesize=A4, title=title, leftMargin=2*cm, rightMargin=2*cm,
                      topMargin=2*cm, bottomMargin=2*cm).build(story or [Paragraph(" ", st["BodyText"])])
    return buf.getvalue()

def render_pptx(blocks, title) -> bytes:
    from pptx import Presentation
    from pptx.util import Pt
    prs = Presentation()
    slides, cur, deck_title = [], None, title
    for b in blocks:
        if b[0] == "h" and b[1] == 1 and not slides and cur is None:
            deck_title = _plain(b[2]); continue
        if b[0] == "h":
            cur = {"title": _plain(b[2]), "lines": []}; slides.append(cur); continue
        if cur is None:
            cur = {"title": deck_title, "lines": []}; slides.append(cur)
        if b[0] in ("p", "ul", "ol"):
            cur["lines"].append(_plain(b[1]))
        elif b[0] == "table":
            cur["lines"] += [" | ".join(_plain(c) for c in r) for r in b[1]]
        elif b[0] == "code":
            cur["lines"] += b[1].splitlines()[:6]

    ts = prs.slides.add_slide(prs.slide_layouts[0])
    ts.shapes.title.text = deck_title
    ts.placeholders[1].text = ""
    for s in slides:
        chunks = [s["lines"][i:i + 6] for i in range(0, len(s["lines"]), 6)] or [[]]
        for k, chunk in enumerate(chunks):
            sl = prs.slides.add_slide(prs.slide_layouts[1])
            sl.shapes.title.text = s["title"] + (" (cont.)" if k else "")
            tf = sl.placeholders[1].text_frame
            for j, line in enumerate(chunk):
                p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
                p.text = line; p.font.size = Pt(20 if len(chunk) <= 5 else 16)
    buf = BytesIO(); prs.save(buf); return buf.getvalue()

def render_xlsx(blocks, title) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook(); wb.remove(wb.active)
    heading, used = "Sheet", set()

    def num(c):
        try:
            return float(c) if re.fullmatch(r"-?\d+\.\d+", c) else int(c)
        except ValueError:
            return _plain(c)

    def sheet_name(base):
        base = re.sub(r"[\[\]:*?/\\]", "", base)[:28] or "Sheet"
        name, k = base, 2
        while name in used:
            name = f"{base}{k}"; k += 1
        used.add(name); return name

    for b in blocks:
        if b[0] == "h":
            heading = _plain(b[2])
        elif b[0] == "table" and b[1]:
            ws = wb.create_sheet(sheet_name(heading))
            for ri, row in enumerate(b[1], 1):
                for ci, c in enumerate(row, 1):
                    cell = ws.cell(ri, ci, c if ri == 1 else num(c))
                    if ri == 1: cell.font = Font(bold=True)
            for col in ws.columns:
                ws.column_dimensions[col[0].column_letter].width = min(
                    max(len(str(c.value or "")) for c in col) + 2, 50)
    if not wb.worksheets:  # model gave no tables -> dump the text
        ws = wb.create_sheet("Sheet1")
        for b in blocks:
            if b[0] in ("p", "ul", "ol"): ws.append([_plain(b[1])])
            elif b[0] == "h": ws.append([_plain(b[2])])
    buf = BytesIO(); wb.save(buf); return buf.getvalue()

RENDERERS = {"pdf": render_pdf, "docx": render_docx, "pptx": render_pptx, "xlsx": render_xlsx}

def build_document(fmt: str, markdown: str) -> tuple[str, bytes]:
    blocks = parse_markdown(markdown)
    first_h = next((b[2] for b in blocks if b[0] == "h"), "document")
    title = _plain(first_h)
    data = RENDERERS[fmt](blocks, title)
    slug = re.sub(r"[^a-z0-9]+", "-", title.lower()).strip("-")[:40] or "document"
    return f"{slug}-{uuid.uuid4().hex[:6]}.{fmt}", data